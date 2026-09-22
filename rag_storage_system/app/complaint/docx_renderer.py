"""DOCX rendering for ComplaintDraft - same python-docx approach as app/report/docx_renderer.py."""

import io

from docx import Document

from app.complaint.schema import ComplaintDraft
from app.complaint.sections import complaint_sections


class ComplaintDocxRenderer:

    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    file_extension = "docx"

    def render(self, complaint: ComplaintDraft) -> bytes:
        document = Document()
        document.add_heading("DRAFT COMPLAINT", level=0)

        for heading, paragraphs in complaint_sections(complaint):
            document.add_heading(heading, level=1)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()