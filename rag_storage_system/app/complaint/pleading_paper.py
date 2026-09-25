"""
Renders a PleadingContent on California pleading paper (California Rules
of Court, rules 2.100-2.119), as a Word document the attorney can edit:

- Letter size. Times New Roman 12 pt (rule 2.104). Left margin 1.5 in and
  right margin 0.5 in (rule 2.107).
- Lines numbered 1-28 down the left margin, with a double rule to the
  right of the numbers and a single rule down the right margin (rule
  2.108). Text is double-spaced on an exact 24 pt grid, so every text line
  sits on a number. Single-spaced blocks (attorney block, caption,
  signature) always use an even number of 12 pt lines, so the grid stays
  aligned after them.
- First page (rule 2.111): attorney name, bar number, address, phone and
  email from line 1; the court's name on line 8; the caption with the
  parties on the left and the case number and title on the right.
- Footer on every page (rule 2.110): page number, a rule, and the title
  of the paper.

The numbers and rules live in the page header, so they are on every page
and stay put while the attorney edits the text.
"""

import io
import textwrap

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from app.complaint.pleading import PleadingContent

_LINE = 480            # 24 pt in twips - one numbered line
_HALF = 240            # 12 pt - a single-spaced line
_LINES_PER_PAGE = 28
_TOP = 1440            # 1 in
_BOTTOM = 15840 - _TOP - _LINES_PER_PAGE * _LINE   # what's left of 11 in: 960 twips
_LEFT = 2160           # 1.5 in
_RIGHT = 720           # 0.5 in
_NUMBERS_X = 1296      # the number column starts 0.9 in from the page edge...
_NUMBERS_WIDTH = 720   # ...and ends (with the double rule) at 1.4 in
_FONT = "Times New Roman"
_CAPTION_WRAP = 36     # characters per caption line (each caption cell is ~3.25 in wide)


# w:pPr children must appear in this order (ECMA-376 CT_PPrBase) - Word and
# LibreOffice ignore or mis-read properties that are out of order.
_PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers",
    "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
    "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
    "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]


def _order_ppr(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    rank = {qn(f"w:{name}"): i for i, name in enumerate(_PPR_ORDER)}
    children = sorted(p_pr, key=lambda child: rank.get(child.tag, len(rank)))
    for child in children:
        p_pr.remove(child)
        p_pr.append(child)


def _set(element, tag: str, **attrs) -> OxmlElement:
    child = element.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        element.append(child)
    for key, value in attrs.items():
        child.set(qn(f"w:{key}"), str(value))
    return child


def _paragraph_format(paragraph, line: int = _LINE, align=None, first_indent: int = 0, left_indent: int = 0,
                      bold: bool = False, size: float = 12) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Twips(line)
    fmt.line_spacing_rule = 4  # WD_LINE_SPACING.EXACTLY
    fmt.widow_control = False
    if align is not None:
        paragraph.alignment = align
    if first_indent:
        fmt.first_line_indent = Twips(first_indent)
    if left_indent:
        fmt.left_indent = Twips(left_indent)
    for run in paragraph.runs:
        run.font.name = _FONT
        run.font.size = Pt(size)
        run.bold = bold


def _add(container, text: str, keep_with_next: bool = False, **fmt) -> None:
    paragraph = container.add_paragraph()
    paragraph.add_run(text)
    _paragraph_format(paragraph, **fmt)
    if keep_with_next:
        paragraph.paragraph_format.keep_with_next = True


def _single_spaced(container, lines: list[str], keep_together: bool = False, **fmt) -> None:
    """Lines at 12 pt, padded to an even count so the 24 pt grid stays aligned after them."""

    if len(lines) % 2:
        lines = lines + [""]
    for index, line in enumerate(lines):
        _add(container, line, line=_HALF, keep_with_next=keep_together and index < len(lines) - 1, **fmt)


def _cell_borders(cell, **sides) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = _set(tc_pr, "w:tcBorders")
    for side, spec in sides.items():
        _set(borders, f"w:{side}", val=spec, sz=6 if spec == "double" else 4, space=0, color="000000")


def _no_cell_margins(table) -> None:  # used by the caption table
    tbl_pr = table._tbl.tblPr
    margins = _set(tbl_pr, "w:tblCellMar")
    for side in ("top", "bottom"):
        _set(margins, f"w:{side}", w=0, type="dxa")


def _setup_page(section) -> None:
    section.page_width = Twips(12240)
    section.page_height = Twips(15840)
    section.top_margin = Twips(_TOP)
    section.bottom_margin = Twips(_BOTTOM)
    section.left_margin = Twips(_LEFT)
    section.right_margin = Twips(_RIGHT)
    section.header_distance = Twips(360)
    section.footer_distance = Twips(160)

    # The single rule down the right margin.
    sect_pr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "text")
    right = OxmlElement("w:right")
    for key, value in {"val": "single", "sz": "4", "space": "4", "color": "000000"}.items():
        right.set(qn(f"w:{key}"), value)
    borders.append(right)
    # pgBorders must come before the column/docGrid settings in sectPr.
    anchor = sect_pr.find(qn("w:cols"))
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        sect_pr.append(borders)


def _line_numbers(section) -> None:
    """
    Numbers 1-28 on the 24 pt grid, in a text frame in the header
    (paragraphs sharing one w:framePr - positioned on the page, so they
    take no room from the text), with the double rule as the frame's
    right border. Frames are understood by Word and LibreOffice alike.
    """

    header = section.header
    header.is_linked_to_previous = False
    for paragraph in list(header.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)

    for number in range(1, _LINES_PER_PAGE + 1):
        paragraph = header.add_paragraph()
        paragraph.add_run(str(number))
        _paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT)
        p_pr = paragraph._p.get_or_add_pPr()
        frame = OxmlElement("w:framePr")
        for key, value in {
            "w": _NUMBERS_WIDTH, "h": _LINES_PER_PAGE * _LINE, "hRule": "exact", "hSpace": 0, "wrap": "around",
            "vAnchor": "page", "hAnchor": "page", "x": _NUMBERS_X, "y": _TOP,
        }.items():
            frame.set(qn(f"w:{key}"), str(value))
        # framePr sits near the start of pPr (after pStyle/keepNext...) - before spacing/jc.
        p_pr.insert(0, frame)
        borders = _set(p_pr, "w:pBdr")
        _set(borders, "w:right", val="double", sz=6, space=4, color="000000")
        paragraph.paragraph_format.right_indent = Twips(120)
        _order_ppr(paragraph)

    # A frame needs an ordinary paragraph after it to anchor to; keep it tiny so the header takes no room.
    anchor = header.add_paragraph()
    _paragraph_format(anchor, line=20, size=1)


def _footer(section, title: str) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    number = footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    number._p.append(field)
    _paragraph_format(number, line=_HALF, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)

    rule = footer.add_paragraph()
    rule.add_run(title.upper())
    _paragraph_format(rule, line=_HALF, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    p_pr = rule._p.get_or_add_pPr()
    borders = _set(p_pr, "w:pBdr")
    _set(borders, "w:top", val="single", sz=6, space=1, color="000000")
    _order_ppr(rule)


def _caption(document, content: PleadingContent) -> None:
    left = [f"{content.plaintiff}, an individual,", "", "Plaintiff,", "", "v.", "",
            f"{content.defendant}; and DOES 1 through 20, inclusive,", "", "Defendants."]
    right = [f"Case No.: {content.case_number}", "", f"{content.title} FOR:"] + content.causes_list

    def wrapped(lines: list[str], indent: str = "") -> list[str]:
        out = []
        for line in lines:
            out += textwrap.wrap(line, _CAPTION_WRAP, subsequent_indent=indent) or [""]
        return out

    left_lines = wrapped(left, "    ")
    right_lines = wrapped(right, "    ")
    rows = max(len(left_lines), len(right_lines))
    rows += rows % 2  # even number of 12 pt lines = whole numbered lines

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _no_cell_margins(table)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    _set(tr_pr, "w:trHeight", val=rows * _HALF, hRule="exact")
    left_cell, right_cell = table.rows[0].cells
    left_cell.width = right_cell.width = Twips(4680)
    _cell_borders(left_cell, right="single", bottom="single")

    for cell, lines, indent in ((left_cell, left_lines, 0), (right_cell, right_lines, 180)):
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        for line in lines:
            paragraph = cell.add_paragraph()
            paragraph.add_run(line)
            _paragraph_format(paragraph, line=_HALF, left_indent=indent)
            # "Plaintiff," / "Defendants." sit indented like a traditional caption.
            if line in ("Plaintiff,", "Defendants."):
                paragraph.paragraph_format.left_indent = Twips(2160)
            if line == "v.":
                paragraph.paragraph_format.left_indent = Twips(1080)


def render_pleading_paper(content: PleadingContent) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = _FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)

    section = document.sections[0]
    _setup_page(section)
    _line_numbers(section)
    _footer(section, content.title)

    # python-docx's default template may start with an empty paragraph - it would shift the grid.
    for paragraph in list(document.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)

    # Lines 1-7: the attorney block (14 single-spaced lines).
    attorney = list(content.attorney_lines)
    attorney += [""] * max(0, 14 - len(attorney))
    _single_spaced(document, attorney)

    # Line 8 on: the court, then the caption.
    for line in content.court_lines:
        _add(document, line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    _add(document, "")
    _caption(document, content)
    _add(document, "")

    signature: list[str] = []
    for position, block in enumerate(content.body):
        # The line before the signature travels with it - a signature alone on a page is not accepted practice.
        before_signature = position + 1 < len(content.body) and content.body[position + 1].kind == "signature"
        if block.kind == "signature":
            signature.append(block.text)
            continue
        if signature:
            _single_spaced(document, signature, keep_together=True, left_indent=5040)
            signature = []
        if block.kind == "heading":
            _add(document, block.text, keep_with_next=True, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        elif block.kind == "subheading":
            _add(document, block.text, keep_with_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif block.kind == "numbered":
            _add(document, f"{block.number}.\t{block.text}", first_indent=720, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        else:
            _add(document, block.text, keep_with_next=before_signature, first_indent=720, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if signature:
        _add(document, "", keep_with_next=True)
        _single_spaced(document, signature, keep_together=True, left_indent=5040)

    if content.attorney_notes:
        page_break = document.add_paragraph()
        page_break.add_run().add_break(WD_BREAK.PAGE)
        _paragraph_format(page_break)
        _add(document, "ATTORNEY NOTES - NOT PART OF THE PLEADING - REMOVE BEFORE FILING",
             align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        for heading, paragraphs in content.attorney_notes:
            _add(document, heading, bold=True)
            for paragraph in paragraphs:
                _add(document, paragraph, first_indent=720)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


LINES_PER_PAGE = _LINES_PER_PAGE
