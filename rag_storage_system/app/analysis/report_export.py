"""
DOCX/PDF export of an analysis report (app/analysis/report_builder.py's
build_analysis_report() dict, same shape as the AnalysisReport schema).

Both formats render the same sections, always ending with the
Owner-editable disclaimer (app/disclaimer.py) - whatever the Owner has
currently saved via PUT /admin/disclaimer, or the built-in default if
they never have, so every export always ships with a disclaimer.

DOCX: python-docx (already a dependency, used elsewhere in this
project only to *read* uploaded .docx files - see
app/extraction/docx_extractor.py). Word paginates a Document's
paragraphs on its own, so no manual page-layout logic is needed here.

PDF: PyMuPDF's Story/DocumentWriter API, which flows HTML across as
many pages as it needs - the same manual-pagination problem word
processors solve for DOCX, solved here by building one HTML string and
letting Story lay it out, rather than positioning text by hand.
"""

import html
import io

import pymupdf
from docx import Document

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MEDIA_TYPE = "application/pdf"


def _source_label(source: dict) -> str:
    label = f"{source['filename']} ({source['category']})"
    if source.get("section"):
        label += f" - section {source['section']}"
    elif source.get("start_page"):
        label += f" - page {source['start_page']}"
    return label


def _report_sections(report: dict, disclaimer_text: str) -> list[tuple[str, list[str]]]:
    """
    Flatten an analysis report dict into (heading, paragraphs) pairs -
    the single content model both build_report_docx() and
    build_report_pdf() render, so the two formats never drift apart.
    """

    breakdown = report["match_score_breakdown"]

    sections: list[tuple[str, list[str]]] = [
        ("Executive Summary", [report["executive_summary"]["text"]]),
        (
            "Overall Match Score",
            [
                f"{report['overall_match_score']:.1f} / 100",
                f"Coverage ratio: {breakdown['coverage_ratio']:.2f}  |  "
                f"Average confidence: {breakdown['avg_confidence']:.2f}",
            ],
        ),
    ]

    recommendations = [item["text"] for item in report["recommendations"]]
    if recommendations:
        sections.append(("Recommendations", recommendations))

    detailed = report["detailed_matching"]
    for label, key in (
        ("Similarities", "similarities"),
        ("Differences", "differences"),
        ("Gaps", "gaps"),
        ("Conflicts", "conflicts"),
    ):
        items = detailed[key]
        if not items:
            continue

        paragraphs = []
        for item in items:
            source_names = ", ".join(_source_label(s) for s in item["sources"]) or "no cited source"
            paragraphs.append(f"{item['narrative']['text']} - Sources: {source_names}")

        sections.append((label, paragraphs))

    sources = [_source_label(s) for s in report["sources"]]
    if sources:
        sections.append(("Sources", sources))

    sections.append(("Disclaimer", [disclaimer_text]))

    return sections


def build_report_docx(report: dict, disclaimer_text: str) -> bytes:
    """Render an analysis report + the current disclaimer as a .docx file."""

    document = Document()
    document.add_heading("Analysis Report", level=0)

    for heading, paragraphs in _report_sections(report, disclaimer_text):
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_report_pdf(report: dict, disclaimer_text: str) -> bytes:
    """Render an analysis report + the current disclaimer as a .pdf file."""

    html_parts = ["<h1>Analysis Report</h1>"]
    for heading, paragraphs in _report_sections(report, disclaimer_text):
        html_parts.append(f"<h2>{html.escape(heading)}</h2>")
        html_parts.extend(f"<p>{html.escape(paragraph)}</p>" for paragraph in paragraphs)

    story = pymupdf.Story(html="".join(html_parts))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)

    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)

    more = 1
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()

    writer.close()
    return buffer.getvalue()
_EXPORT_EXCERPT_CHARS = 700


def _source_with_excerpt(source: dict) -> str:
    """The source label, then the passage it was cited for - quoted, so the memo is verifiable on its own."""

    label = _source_label(source)
    excerpt = " ".join((source.get("excerpt") or "").split())
    if not excerpt:
        return label
    if len(excerpt) > _EXPORT_EXCERPT_CHARS:
        excerpt = excerpt[:_EXPORT_EXCERPT_CHARS].rsplit(" ", 1)[0] + " ..."
    return f"{label}\n\u201c{excerpt}\u201d"


def _memo_blocks(
    title: str, exchanges: list[dict], disclaimer_text: str, prepared_by: str | None, prepared_on: str
) -> list[tuple[str, str, list[str]]]:
    """
    A research memorandum (Blueprint Phase 2: "heading, question presented,
    analysis, authorities cited") as (kind, heading, paragraphs) blocks -
    the one content model both DOCX and PDF render, so they can't drift.
    `exchanges` are {"question", "answer", "sources"} exactly as shown on
    screen / saved in the thread; nothing is re-retrieved or re-generated.
    """

    header = [f"RE: {title}", f"DATE: {prepared_on}"]
    if prepared_by:
        header.append(f"FROM: {prepared_by}")
    blocks: list[tuple[str, str, list[str]]] = [("header", "", header)]

    numbered = len(exchanges) > 1
    for number, exchange in enumerate(exchanges, start=1):
        suffix = f" {number}" if numbered else ""
        sources = exchange.get("sources") or []
        blocks.append(("h1", f"Question Presented{suffix}", [exchange["question"]]))
        blocks.append(("h2", "Analysis", [exchange["answer"]]))
        blocks.append((
            "h2", "Authorities Cited",
            [_source_with_excerpt(s) for s in sources]
            or ["None - the firm's library did not contain authority on this point."],
        ))

    blocks.append(("h1", "Disclaimer", [disclaimer_text]))
    return blocks


def build_research_memo_docx(
    title: str, exchanges: list[dict], disclaimer_text: str, prepared_by: str | None = None
) -> bytes:
    """Render research Q&A (one answer or a whole thread) as a .docx memorandum."""

    from datetime import date

    document = Document()
    document.add_heading("Research Memorandum", level=0)

    for kind, heading, paragraphs in _memo_blocks(title, exchanges, disclaimer_text, prepared_by, date.today().isoformat()):
        if kind == "h1":
            document.add_heading(heading, level=1)
        elif kind == "h2":
            document.add_heading(heading, level=2)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_research_memo_pdf(
    title: str, exchanges: list[dict], disclaimer_text: str, prepared_by: str | None = None
) -> bytes:
    """Render research Q&A (one answer or a whole thread) as a .pdf memorandum."""

    from datetime import date

    html_parts = ["<h1>Research Memorandum</h1>"]
    for kind, heading, paragraphs in _memo_blocks(title, exchanges, disclaimer_text, prepared_by, date.today().isoformat()):
        if kind in ("h1", "h2"):
            html_parts.append(f"<{kind}>{html.escape(heading)}</{kind}>")
        html_parts.extend(f"<p>{html.escape(paragraph).replace(chr(10), '<br/>')}</p>" for paragraph in paragraphs)

    story = pymupdf.Story(html="".join(html_parts))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)

    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)

    more = 1
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()

    writer.close()
    return buffer.getvalue()


def build_owner_research_docx(query: str, answer: str, sources: list[dict], disclaimer_text: str) -> bytes:
    """A single Owner Research answer as a memo (see build_research_memo_docx())."""

    return build_research_memo_docx(query, [{"question": query, "answer": answer, "sources": sources}], disclaimer_text)


def build_owner_research_pdf(query: str, answer: str, sources: list[dict], disclaimer_text: str) -> bytes:
    """A single Owner Research answer as a memo (see build_research_memo_pdf())."""

    return build_research_memo_pdf(query, [{"question": query, "answer": answer, "sources": sources}], disclaimer_text)
