"""DOCX rendering for StructuredReport - same python-docx approach as app/analysis/report_export.py's build_report_docx()."""

import io

from docx import Document

from app.report.renderer_base import ReportRenderer
from app.report.schema import StructuredReport
from app.report.sections import report_sections


class DocxReportRenderer(ReportRenderer):

    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    file_extension = "docx"

    def render(self, report: StructuredReport) -> bytes:
        document = Document()
        document.add_heading("Client Intake Report", level=0)

        for heading, paragraphs in report_sections(report):
            document.add_heading(heading, level=1)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
