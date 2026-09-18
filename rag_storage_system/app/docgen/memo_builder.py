"""
The research memo, as .docx and .pdf.

Input is one dict, the memo spec:

    {
      "title":        str,
      "prepared_for": str,
      "generated_at": str,
      "disclaimer":   str,
      "entries": [ {"question": str,
                    "analysis": str,
                    "sources": [ {marker, filename, category, chapter,
                                  section, start_page, end_page, score} ]} ],
    }

Structure per Blueprint Phase 2 step 4: heading, question presented,
analysis, authorities cited.

.docx via python-docx (already a dependency - used today by
app/extraction/docx_extractor.py to read). .pdf via pymupdf's Story
(also already a dependency - used today by app/extraction/pdf_extractor.py),
so nothing new is added to requirements.txt.

When the owner supplies formatting exemplars, this module is where a
real .docx template goes: load the template with Document(path)
instead of Document() and write into its styles. The function
signatures do not change.
"""

import io
from datetime import datetime, timezone

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _format_location(source: dict) -> str:
    """'filename - Chapter > Section (pp. 3-4)' - traceable at a glance."""

    parts = [source.get("filename") or "unknown"]

    location = " > ".join(
        str(source[key]) for key in ("chapter", "section") if source.get(key)
    )
    if location:
        parts.append(location)

    if source.get("start_page") is not None:
        end = source.get("end_page", source["start_page"])
        parts.append(f"pp. {source['start_page']}-{end}")

    return " - ".join(parts[:2]) + (f" ({parts[2]})" if len(parts) > 2 else "")


def default_memo(title: str, prepared_for: str, entries: list[dict],
                 disclaimer: str) -> dict:
    return {
        "title": title,
        "prepared_for": prepared_for,
        "generated_at": datetime.now(timezone.utc).strftime("%d %B %Y"),
        "disclaimer": disclaimer,
        "entries": entries,
    }


# ----------------------------------------------------------------------
# .docx
# ----------------------------------------------------------------------

def build_memo_docx(memo: dict) -> bytes:
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    heading = document.add_heading(memo["title"], level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"Prepared for {memo['prepared_for']}  |  {memo['generated_at']}"
    )
    meta_run.italic = True

    for index, entry in enumerate(memo["entries"], start=1):
        document.add_heading(f"{index}. Question Presented", level=1)
        document.add_paragraph(entry["question"])

        document.add_heading("Analysis", level=1)
        for block in str(entry["analysis"]).split("\n\n"):
            if block.strip():
                document.add_paragraph(block.strip())

        document.add_heading("Authorities Cited", level=1)

        if entry.get("sources"):
            for source in entry["sources"]:
                bullet = document.add_paragraph(style="List Bullet")
                bullet.add_run(f"[{source.get('marker', '?')}] ").bold = True
                bullet.add_run(_format_location(source))
        else:
            none_cited = document.add_paragraph()
            none_cited.add_run(
                "No library authority was cited for this question."
            ).italic = True

        if index < len(memo["entries"]):
            document.add_page_break()

    document.add_paragraph()
    footer = document.add_paragraph()
    footer.add_run(memo["disclaimer"]).italic = True

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


# ----------------------------------------------------------------------
# .pdf
# ----------------------------------------------------------------------

_PDF_CSS = """
body   { font-family: sans-serif; font-size: 10.5pt; color: #111; }
h1     { font-size: 17pt; text-align: center; margin-bottom: 2pt; }
.meta  { text-align: center; font-style: italic; color: #555;
         font-size: 9.5pt; margin-bottom: 14pt; }
h2     { font-size: 12pt; color: #1a3a5c; margin-top: 14pt;
         margin-bottom: 4pt; }
p      { margin: 0 0 7pt 0; line-height: 1.45; }
li     { margin-bottom: 3pt; }
.marker{ font-weight: bold; }
.disc  { font-style: italic; color: #555; font-size: 9pt;
         margin-top: 18pt; border-top: 1px solid #ccc; padding-top: 6pt; }
"""


def _escape(text: str) -> str:
    return (
        str(text)
        .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


def _memo_html(memo: dict) -> str:
    parts = [
        f"<h1>{_escape(memo['title'])}</h1>",
        f"<p class='meta'>Prepared for {_escape(memo['prepared_for'])} "
        f"&nbsp;|&nbsp; {_escape(memo['generated_at'])}</p>",
    ]

    for index, entry in enumerate(memo["entries"], start=1):
        parts.append(f"<h2>{index}. Question Presented</h2>")
        parts.append(f"<p>{_escape(entry['question'])}</p>")
        parts.append("<h2>Analysis</h2>")

        for block in str(entry["analysis"]).split("\n\n"):
            if block.strip():
                parts.append(f"<p>{_escape(block.strip())}</p>")

        parts.append("<h2>Authorities Cited</h2>")

        if entry.get("sources"):
            parts.append("<ul>")
            for source in entry["sources"]:
                parts.append(
                    f"<li><span class='marker'>[{_escape(source.get('marker', '?'))}]</span> "
                    f"{_escape(_format_location(source))}</li>"
                )
            parts.append("</ul>")
        else:
            parts.append("<p><i>No library authority was cited for this question.</i></p>")

    parts.append(f"<p class='disc'>{_escape(memo['disclaimer'])}</p>")

    return f"<html><body>{''.join(parts)}</body></html>"


def build_memo_pdf(memo: dict) -> bytes:
    """
    Render via pymupdf's Story engine: it paginates the flowed HTML
    itself, so a long answer becomes a multi-page memo rather than a
    clipped one.
    """

    story = pymupdf.Story(html=_memo_html(memo), user_css=_PDF_CSS)

    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)

    page_rect = pymupdf.paper_rect("letter")
    content_rect = page_rect + (54, 54, -54, -54)   # 0.75in margins

    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()

    writer.close()

    return buffer.getvalue()