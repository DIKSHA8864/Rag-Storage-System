"""
California pleading paper (app/complaint/pleading_paper.py), the firm's
own .docx templates (app/complaint/template_fill.py), and their API
(app/api/pleading_api.py + the complaint generator's `style`).
"""

import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from fastapi.testclient import TestClient

from app.api import storage_api
from app.complaint.pleading import build_pleading_content
from app.complaint.pleading_paper import render_pleading_paper
from app.complaint.schema import ComplaintCauseOfAction, ComplaintDraft, ComplaintElement, ResearchSuggestion
from app.complaint.template_fill import TemplateError, fill_template, inspect_template
from app.metadata.sqlite_repository import SQLiteMetadataRepository
from app.security.auth import require_admin_key
from app.storage.local_backend import LocalStorageBackend

SETTINGS = {
    "attorney_name": "Jane Q. Attorney", "bar_number": "123456", "firm_name": "AshiLegal LLP",
    "address": "100 Main Street\nLos Angeles, CA 90012", "phone": "(213) 555-0100", "email": "jane@example.com",
    "attorney_for": "Plaintiff", "county": "Los Angeles",
}


def _draft() -> ComplaintDraft:
    return ComplaintDraft(
        intake_session_id=1, matter_name="Case", generated_at="now", plaintiff_name="Case",
        causes_of_action=[
            ComplaintCauseOfAction(
                cause_of_action_id=1, name="Failure to Pay Overtime", authority_citation="Cal. Lab. Code § 510",
                elements=[
                    ComplaintElement(description="Plaintiff worked overtime", satisfied_by="About 55 hours a week."),
                    ComplaintElement(description="Overtime was not paid", placeholder="[ATTORNEY TO PROVIDE: pay records]"),
                ],
                research_suggestions=[ResearchSuggestion(filename="wage.pdf", category="Wage", chunk_text="RESEARCH TEXT", score=0.5)],
            ),
            ComplaintCauseOfAction(
                cause_of_action_id=2, name="Retaliation", authority_citation="Cal. Lab. Code § 98.6",
                elements=[ComplaintElement(description="Protected activity", satisfied_by="Complained to HR.")],
            ),
        ],
        disclaimer_text="Not legal advice.",
    )


# ----------------------------------------------------------------------
# Content
# ----------------------------------------------------------------------

def test_unknown_details_stay_bracketed_placeholders():
    content = build_pleading_content(_draft(), None)

    assert content.plaintiff == "[PLAINTIFF NAME]"
    assert content.defendant == "[DEFENDANT NAME]"
    assert content.case_number == "[CASE NUMBER]"
    assert content.attorney_lines[0] == "[ATTORNEY NAME] (SBN [STATE BAR NO.])"
    assert content.court_lines == ["SUPERIOR COURT OF THE STATE OF CALIFORNIA", "FOR THE COUNTY OF [COUNTY]"]


def test_allegations_are_numbered_in_order_and_research_stays_off_the_pleading():
    content = build_pleading_content(_draft(), SETTINGS, plaintiff="JOHN DOE", defendant="ACME, INC.", case_number="23STCV01")

    numbers = [b.number for b in content.body if b.kind == "numbered"]
    assert numbers == list(range(1, len(numbers) + 1))
    headings = [b.text for b in content.body if b.kind == "heading"]
    assert headings == ["PARTIES", "JURISDICTION AND VENUE", "FIRST CAUSE OF ACTION", "SECOND CAUSE OF ACTION", "PRAYER FOR RELIEF"]
    body_text = " ".join(b.text for b in content.body)
    assert "About 55 hours a week." in body_text and "[ATTORNEY TO PROVIDE: pay records]" in body_text
    assert "Authority: Cal. Lab. Code § 510" in body_text
    assert "RESEARCH TEXT" not in body_text
    assert any("RESEARCH TEXT" in " ".join(paragraphs) for _, paragraphs in content.attorney_notes)
    assert content.causes_list == ["1. FAILURE TO PAY OVERTIME", "2. RETALIATION"]
    assert content.court_lines[1] == "FOR THE COUNTY OF LOS ANGELES"


# ----------------------------------------------------------------------
# Pleading paper
# ----------------------------------------------------------------------

def test_pleading_paper_layout_follows_the_rules_of_court():
    document = Document(io.BytesIO(render_pleading_paper(build_pleading_content(_draft(), SETTINGS, plaintiff="JOHN DOE"))))
    section = document.sections[0]

    assert (section.page_width.inches, section.page_height.inches) == (8.5, 11)
    assert section.left_margin.inches == 1.5 and section.right_margin.inches == 0.5

    numbers = [p for p in section.header.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:framePr")) is not None]
    assert [p.text for p in numbers] == [str(n) for n in range(1, 29)]
    assert all(p.paragraph_format.line_spacing == Pt(24) for p in numbers)
    assert numbers[0]._p.pPr.find(qn("w:pBdr")).find(qn("w:right")).get(qn("w:val")) == "double"
    assert section._sectPr.find(qn("w:pgBorders")) is not None  # the rule down the right margin

    footer_xml = section.footer._element.xml
    assert 'w:instr="PAGE"' in footer_xml
    assert "COMPLAINT FOR DAMAGES" in section.footer.paragraphs[-1].text

    body = [p for p in document.paragraphs if p.text.strip()]
    assert body[0].text == "Jane Q. Attorney (SBN 123456)"
    assert "SUPERIOR COURT OF THE STATE OF CALIFORNIA" in [p.text for p in body]
    assert all(p.paragraph_format.line_spacing in (Pt(12), Pt(24)) for p in document.paragraphs)
    caption = document.tables[0].rows[0].cells
    assert "JOHN DOE, an individual," in caption[0].text and "Case No.: [CASE NUMBER]" in caption[1].text
    text = "\n".join(p.text for p in document.paragraphs)
    assert "1.\tPlaintiff JOHN DOE is, and at all relevant times was, an individual." in text
    assert "ATTORNEY NOTES - NOT PART OF THE PLEADING - REMOVE BEFORE FILING" in text


def test_single_spaced_blocks_keep_the_24pt_grid():
    """Every run of 12 pt lines is an even count, so the double-spaced text after it still sits on a number."""

    document = Document(io.BytesIO(render_pleading_paper(build_pleading_content(_draft(), SETTINGS))))
    run = 0
    for paragraph in document.paragraphs:
        if paragraph.paragraph_format.line_spacing == Pt(12):
            run += 1
        else:
            assert run % 2 == 0
            run = 0


# ----------------------------------------------------------------------
# The firm's own templates
# ----------------------------------------------------------------------

def _template(body_line="{{body}}", extra=None) -> bytes:
    document = Document()
    header = document.sections[0].header.paragraphs[0]
    header.add_run("{{attorney_")
    header.add_run("block}}")  # split across runs, as Word often does
    title = document.add_paragraph()
    title.add_run("{{court}}").bold = True
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "{{plaintiff}} v. {{defendant}}"
    table.rows[0].cells[1].text = "Case No. {{case_number}} - {{title}}"
    body = document.add_paragraph()
    run = body.add_run(body_line)
    run.font.name = "Courier New"
    if extra:
        document.add_paragraph(extra)
    document.add_paragraph("END OF TEMPLATE")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_template_is_inspected_before_use():
    assert inspect_template(_template()) == ["attorney_block", "body", "case_number", "court", "defendant", "plaintiff", "title"]
    with pytest.raises(TemplateError, match="plaintif_name"):
        inspect_template(_template(extra="{{plaintif_name}}"))
    with pytest.raises(TemplateError, match="needs a"):
        inspect_template(_template(body_line="no body here"))
    with pytest.raises(TemplateError, match="own"):
        inspect_template(_template(body_line="Text before {{body}}"))
    with pytest.raises(TemplateError, match="Word"):
        inspect_template(b"not a docx")


def test_fill_template_places_everything_and_keeps_the_firms_formatting():
    content = build_pleading_content(_draft(), SETTINGS, plaintiff="JOHN DOE", defendant="ACME, INC.", case_number="23STCV01")
    document = Document(io.BytesIO(fill_template(_template(), content)))

    assert document.sections[0].header.paragraphs[0].text.startswith("Jane Q. Attorney (SBN 123456)\nAshiLegal LLP")
    assert document.tables[0].rows[0].cells[0].text == "JOHN DOE v. ACME, INC."
    assert document.tables[0].rows[0].cells[1].text == "Case No. 23STCV01 - COMPLAINT FOR DAMAGES"
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "SUPERIOR COURT OF THE STATE OF CALIFORNIA\nFOR THE COUNTY OF LOS ANGELES"
    assert texts.index("Plaintiff JOHN DOE alleges as follows:") < texts.index("PARTIES") < texts.index("END OF TEMPLATE")
    assert not any("{{" in t for t in texts)
    inserted = document.paragraphs[texts.index("PARTIES")]
    assert inserted.runs[0].font.name == "Courier New"  # the {{body}} paragraph's formatting


# ----------------------------------------------------------------------
# API
# ----------------------------------------------------------------------

@pytest.fixture
def repo(tmp_path):
    return SQLiteMetadataRepository(tmp_path / "metadata.db")


@pytest.fixture
def client(tmp_path, monkeypatch, repo):
    import app.storage as storage_module
    from app.complaint import builder as builder_module

    monkeypatch.setattr(storage_api, "metadata_repository", repo)
    backend = LocalStorageBackend(originals_dir=tmp_path / "intake", quarantine_dir=tmp_path / "intake_quarantine")
    monkeypatch.setattr(storage_module, "get_intake_storage_backend", lambda: backend)
    monkeypatch.setattr(builder_module, "retrieve_for_matter", lambda *a, **k: [])
    yield TestClient(storage_api.app)
    storage_api.app.dependency_overrides.pop(require_admin_key, None)


def _as(role="owner", tenant_id=1):
    storage_api.app.dependency_overrides[require_admin_key] = lambda: {
        "sub": "5", "email": f"{role}@example.com", "role": role, "tenant_id": tenant_id,
    }


def test_pleading_settings_are_per_organization_and_owner_edited(client, repo):
    assert client.get("/admin/pleading-settings").json()["attorney_name"] == ""

    saved = client.put("/admin/pleading-settings", json={**SETTINGS, "attorney_name": "  Jane Q. Attorney  "})
    assert saved.status_code == 200
    assert saved.json()["attorney_name"] == "Jane Q. Attorney"
    assert saved.json()["updated_by"] == "test-owner@example.com"

    _as("attorney")
    assert client.get("/admin/pleading-settings").json()["bar_number"] == "123456"
    assert client.put("/admin/pleading-settings", json=SETTINGS).status_code == 403
    _as("owner", tenant_id=2)
    assert client.get("/admin/pleading-settings").json()["attorney_name"] == ""


def _upload_template(client, data=None, name="Firm pleading"):
    return client.post(
        "/admin/templates", data={"name": name},
        files={"file": ("firm.docx", io.BytesIO(data or _template()), "application/octet-stream")},
    )


def test_templates_upload_list_download_delete(client):
    uploaded = _upload_template(client)
    assert uploaded.status_code == 200
    template = uploaded.json()
    assert "body" in template["placeholders"]

    listed = client.get("/admin/templates").json()
    assert [t["id"] for t in listed["templates"]] == [template["id"]]
    assert "case_number" in listed["placeholders"]
    assert client.get(f"/admin/templates/{template['id']}/download").content == _template()

    bad = _upload_template(client, data=_template(extra="{{typo}}"))
    assert bad.status_code == 400 and "{{typo}}" in bad.json()["detail"]
    assert client.post("/admin/templates", data={"name": "x"},
                       files={"file": ("x.pdf", io.BytesIO(b"%PDF"), "application/pdf")}).status_code == 400

    _as("owner", tenant_id=2)
    assert client.get("/admin/templates").json()["templates"] == []
    assert client.get(f"/admin/templates/{template['id']}/download").status_code == 404
    assert client.delete(f"/admin/templates/{template['id']}").status_code == 404
    _as("attorney")
    assert _upload_template(client).status_code == 403
    _as("owner")

    assert client.delete(f"/admin/templates/{template['id']}").json() == {"deleted": True}
    assert client.get("/admin/templates").json()["templates"] == []


def _session_with_cause(repo):
    session = repo.create_intake_session(matter_id=0, title="Intake")
    repo.add_intake_fact(session["id"], "mandatory_sweep", "overtime", "I worked 50 hours a week without overtime pay.")
    cause = repo.create_cause_of_action("Wage", "Failure to Pay Overtime", ["Employee worked overtime hours without overtime pay"], "Labor Code Section 510")
    return session["id"], cause["id"]


def _generate(client, session_id, **body):
    response = client.post(f"/admin/intake/sessions/{session_id}/complaint", json=body)
    assert response.status_code == 200, response.text
    download = client.get(f"/admin/complaints/{response.json()['id']}/download")
    return Document(io.BytesIO(download.content))


def test_complaint_styles(client, repo):
    session_id, cause_id = _session_with_cause(repo)
    client.put("/admin/pleading-settings", json=SETTINGS)

    pleading = _generate(client, session_id, cause_of_action_ids=[cause_id], plaintiff_name="JOHN DOE", case_number="23STCV01")
    assert pleading.paragraphs[0].text == "Jane Q. Attorney (SBN 123456)"
    assert "Case No.: 23STCV01" in pleading.tables[0].rows[0].cells[1].text
    assert len([p for p in pleading.sections[0].header.paragraphs if p.text.isdigit()]) == 28

    plain = _generate(client, session_id, cause_of_action_ids=[cause_id], style="plain")
    assert plain.paragraphs[0].text == "DRAFT COMPLAINT"

    template_id = _upload_template(client).json()["id"]
    filled = _generate(client, session_id, cause_of_action_ids=[cause_id], style="template", template_id=template_id,
                       defendant_name="ACME, INC.")
    assert "ACME, INC." in filled.tables[0].rows[0].cells[0].text
    assert "FIRST CAUSE OF ACTION" in [p.text for p in filled.paragraphs]

    events = [e["description"] for e in repo.list_timeline_events(session_id) if e["event_type"] == "complaint_generated"]
    assert [e.split(")")[0] for e in events] == [
        "Draft complaint (California pleading paper", "Draft complaint (plain draft", "Draft complaint (template 'Firm pleading'",
    ]

    assert client.post(f"/admin/intake/sessions/{session_id}/complaint",
                       json={"cause_of_action_ids": [cause_id], "style": "fancy"}).status_code == 400
    assert client.post(f"/admin/intake/sessions/{session_id}/complaint",
                       json={"cause_of_action_ids": [cause_id], "style": "template", "template_id": 999}).status_code == 404
