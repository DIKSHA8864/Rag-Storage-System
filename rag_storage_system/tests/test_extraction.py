"""
Extraction tests (Phase 3), porting the manual checks
scripts/test_extraction.py used to only print, into real assertions.
"""

import json

import docx
import pymupdf
import pytest

from app.extraction import extractor_manager
from app.extraction.docx_extractor import extract_docx
from app.extraction.extractor_manager import extract_all_documents, extract_document
from app.extraction.pdf_extractor import extract_pdf
from app.extraction.txt_extractor import extract_txt


# ---------------------------------------------------------------------
# Fixtures - build minimal real PDF/DOCX/TXT files (no binary fixtures
# checked into the repo; pymupdf and python-docx can both author
# files, not just read them).
# ---------------------------------------------------------------------


@pytest.fixture
def sample_pdf(tmp_path):
    path = tmp_path / "sample.pdf"

    document = pymupdf.open()
    for text in ["First page text.", "Second page text."]:
        page = document.new_page()
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()

    return path


@pytest.fixture
def sample_docx(tmp_path):
    path = tmp_path / "sample.docx"

    document = docx.Document()
    document.add_paragraph("Heading paragraph", style="Heading 1")
    document.add_paragraph("Body paragraph one.")
    document.add_paragraph("Body paragraph two.")
    document.save(path)

    return path


@pytest.fixture
def sample_txt(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("First block.\n\nSecond block.\n", encoding="utf-8")
    return path


# ---------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------


def test_extract_pdf_preserves_page_boundaries(sample_pdf):
    result = extract_pdf(sample_pdf)

    assert result["file_type"] == "pdf"
    assert result["page_count"] == 2
    assert [p["page_number"] for p in result["pages"]] == [1, 2]
    assert "First page text" in result["pages"][0]["text"]
    assert "Second page text" in result["pages"][1]["text"]


def test_extract_pdf_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        extract_pdf(tmp_path / "missing.pdf")


@pytest.fixture
def scanned_pdf(tmp_path):
    """A PDF page with no text layer at all - stands in for a scanned page."""

    path = tmp_path / "scanned.pdf"

    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()

    return path


def test_extract_pdf_falls_back_to_ocr_for_a_page_with_no_text_layer(scanned_pdf):
    result = extract_pdf(scanned_pdf)

    assert result["ocr_pages_used"] == 1
    page = result["pages"][0]
    assert page["ocr_used"] is True
    assert page["ocr_is_mock"] is True
    # The mock OCR provider (app/multimodal/ocr.py) never fabricates
    # plausible-looking text - it returns a clearly labeled placeholder,
    # so this proves the fallback actually ran rather than silently
    # leaving the page empty.
    assert "MOCK OCR" in page["text"]


def test_extract_pdf_does_not_run_ocr_when_real_text_extracts_fine(sample_pdf):
    result = extract_pdf(sample_pdf)

    assert result["ocr_pages_used"] == 0
    assert all("ocr_used" not in page for page in result["pages"])


def test_extract_docx_skips_blank_paragraphs_and_keeps_style(sample_docx):
    result = extract_docx(sample_docx)

    assert result["file_type"] == "docx"
    assert result["paragraph_count"] == 3
    assert result["paragraphs"][0]["style"] == "Heading 1"
    assert result["paragraphs"][0]["text"] == "Heading paragraph"
    # Synthetic "pages" mirror paragraphs 1:1 for downstream segmentation.
    assert result["page_count"] == 3


def test_extract_txt_splits_on_blank_lines(sample_txt):
    result = extract_txt(sample_txt)

    assert result["file_type"] == "txt"
    assert result["page_count"] == 2
    assert result["pages"][0]["text"] == "First block."
    assert result["pages"][1]["text"] == "Second block."


def test_extract_txt_single_block_when_no_blank_lines(tmp_path):
    path = tmp_path / "one_block.txt"
    path.write_text("Just one paragraph, no blank lines.", encoding="utf-8")

    result = extract_txt(path)

    assert result["page_count"] == 1


# ---------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------


def test_extract_document_dispatches_by_extension(sample_txt, sample_docx, sample_pdf):
    assert extract_document(sample_txt)["file_type"] == "txt"
    assert extract_document(sample_docx)["file_type"] == "docx"
    assert extract_document(sample_pdf)["file_type"] == "pdf"


def test_extract_document_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "file.md"
    path.write_text("hello", encoding="utf-8")

    with pytest.raises(ValueError):
        extract_document(path)


def test_extract_document_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        extract_document(tmp_path / "missing.txt")


# ---------------------------------------------------------------------
# extract_all_documents() - mirrors storage/originals -> storage/processed
# ---------------------------------------------------------------------


def test_extract_all_documents_mirrors_category_structure(tmp_path, monkeypatch):
    originals = tmp_path / "originals"
    processed = tmp_path / "processed"

    monkeypatch.setattr(extractor_manager, "ORIGINALS_DIR", originals)
    monkeypatch.setattr(extractor_manager, "PROCESSED_DIR", processed)

    category_dir = originals / "Contracts" / "2024"
    category_dir.mkdir(parents=True)
    (category_dir / "note.txt").write_text("Some content.", encoding="utf-8")

    results = extract_all_documents()

    assert len(results) == 1
    assert results[0]["status"] == "extracted"
    assert results[0]["category"] == "Contracts/2024"

    output_path = processed / "Contracts" / "2024" / "note" / "extracted.json"
    assert output_path.exists()

    with open(output_path, encoding="utf-8") as f:
        data = json.load(f)
    assert data["file_type"] == "txt"


def test_extract_all_documents_records_failures(tmp_path, monkeypatch):
    originals = tmp_path / "originals"
    processed = tmp_path / "processed"

    monkeypatch.setattr(extractor_manager, "ORIGINALS_DIR", originals)
    monkeypatch.setattr(extractor_manager, "PROCESSED_DIR", processed)

    # A .pdf extension with non-PDF bytes fails inside pymupdf.
    bad_dir = originals / "Bad"
    bad_dir.mkdir(parents=True)
    (bad_dir / "broken.pdf").write_text("not a real pdf", encoding="utf-8")

    results = extract_all_documents()

    assert len(results) == 1
    assert results[0]["status"] == "failed"
    assert results[0]["category"] == "Bad"


def test_extract_all_documents_empty_when_no_originals(tmp_path, monkeypatch):
    monkeypatch.setattr(extractor_manager, "ORIGINALS_DIR", tmp_path / "does_not_exist")
    assert extract_all_documents() == []
